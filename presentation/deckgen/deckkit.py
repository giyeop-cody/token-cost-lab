# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import copy

BG    = RGBColor(0x0E,0x14,0x1B)
CARD  = RGBColor(0x18,0x22,0x2E)
CARD2 = RGBColor(0x11,0x1A,0x24)
FG    = RGBColor(0xEC,0xF1,0xF6)
MUTED = RGBColor(0x93,0xA4,0xB5)
ACC   = RGBColor(0x36,0xD3,0x99)   # green
ACC2  = RGBColor(0x5A,0x9C,0xFF)  # blue
WARN  = RGBColor(0xFF,0xB0,0x4D)
RED   = RGBColor(0xFF,0x6B,0x6B)
LINE  = RGBColor(0x2A,0x38,0x48)

KFONT = "맑은 고딕"
W = 13.333; H = 7.5

prs = None
BLANK = None

def new_deck():
    """새 프레젠테이션을 만들고 모듈 전역(prs/BLANK/_SLIDE_N)을 초기화한다."""
    global prs, BLANK
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]
    _SLIDE_N[0] = 1
    return prs

def blank_slide():
    """표지처럼 slide() 헬퍼를 쓰지 않는 장.

    카운터는 건드리지 않는다. _SLIDE_N은 "지금까지 부여된 마지막 번호"를
    뜻하고 초기값이 1(=표지)이므로, 표지를 만들 때 올리면 다음 장이 3이 되어
    2번이 통째로 빈다.
    """
    return prs.slides.add_slide(BLANK)

def save(path):
    import os
    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)
    print('saved: %s (%d slides)' % (path, len(prs.slides._sldIdLst)))
    return path

def setfont(run, size, bold=False, color=FG, font=KFONT):
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin','a:ea','a:cs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set('typeface', font)

def rect(sl, x,y,w,h, fill=None, line=None, lw=1.25, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x),Inches(y),Inches(w),Inches(h))
    if radius:
        try: shp.adjustments[0] = radius
        except Exception: pass
    if fill: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    else: shp.fill.background()
    if line: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    else: shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def text(sl, x,y,w,h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tb.text_frame.vertical_anchor = anchor
    first = True
    for p in parts:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', align)
        para.line_spacing = p.get('ls', spacing)
        if 'space_before' in p: para.space_before = Pt(p['space_before'])
        r = para.add_run(); r.text = p['t']
        setfont(r, p.get('sz',18), p.get('b',False), p.get('c',FG), p.get('f',KFONT))
    return tb

_SLIDE_N = [1]  # 표지가 1번이므로 slide()가 만드는 첫 장은 2번

def slide(title=None, kicker=None, n=None):
    _SLIDE_N[0] += 1
    n = _SLIDE_N[0]
    sl = prs.slides.add_slide(BLANK)
    rect(sl,0,0,W,H, fill=BG)
    rect(sl,0,0,0.13,H, fill=ACC)
    if kicker:
        text(sl,0.75,0.42,10,0.32,[{'t':kicker,'sz':13,'b':True,'c':ACC}])
    if title:
        text(sl,0.75,0.75,11.9,0.9,[{'t':title,'sz':32,'b':True,'c':FG}])
        rect(sl,0.75,1.72,1.5,0.045, fill=ACC)
    if n is not None:
        text(sl,12.2,6.95,0.9,0.3,[{'t':str(n),'sz':11,'c':MUTED,'align':PP_ALIGN.RIGHT}])
    return sl

def bullets(sl, x,y,w, items, sz=17, gap=0.52, dot=ACC):
    for i,(head,body) in enumerate(items):
        yy = y + i*gap
        rect(sl,x,yy+0.11,0.12,0.12, fill=dot)
        parts=[{'t':head,'sz':sz,'b':True,'c':FG}]
        tb=text(sl,x+0.32,yy,w-0.32,0.4,parts)
        if body:
            text(sl,x+0.32,yy+0.3,w-0.32,0.4,[{'t':body,'sz':13.5,'c':MUTED}])

def table(sl, x,y,w, cols, rows, colw=None, sz=13, hsz=13, head_c=ACC, zebra=True, rowh=0.42):
    ncol=len(cols)
    colw = colw or [w/ncol]*ncol
    # header
    cx=x
    rect(sl,x,y,w,rowh, fill=CARD)
    for j,c in enumerate(cols):
        text(sl,cx+0.14,y+0.09,colw[j]-0.2,rowh,[{'t':c,'sz':hsz,'b':True,'c':head_c}])
        cx+=colw[j]
    for i,r in enumerate(rows):
        yy=y+rowh+i*rowh
        if zebra and i%2==0: rect(sl,x,yy,w,rowh, fill=CARD2)
        cx=x
        for j,cell in enumerate(r):
            col = FG; bold=False
            if isinstance(cell,tuple): cell,col = cell[0],cell[1]; bold=True
            text(sl,cx+0.14,yy+0.09,colw[j]-0.2,rowh,[{'t':str(cell),'sz':sz,'b':bold,'c':col}])
            cx+=colw[j]
    return y+rowh*(len(rows)+1)

def card(sl,x,y,w,h,title,lines,accent=ACC,tsz=19,bsz=14):
    rect(sl,x,y,w,h, fill=CARD, radius=0.06)
    rect(sl,x,y,0.055,h, fill=accent)
    text(sl,x+0.34,y+0.26,w-0.6,0.4,[{'t':title,'sz':tsz,'b':True,'c':FG}])
    parts=[]
    for i,l in enumerate(lines):
        parts.append({'t':l,'sz':bsz,'c':MUTED,'ls':1.35,'space_before':6 if i else 0})
    text(sl,x+0.34,y+0.85,w-0.6,h-1.0,parts)

def bigstat(sl,x,y,w,val,label,color=ACC,vsz=44):
    rect(sl,x,y,w,1.55, fill=CARD, radius=0.08)
    text(sl,x,y+0.24,w,0.7,[{'t':val,'sz':vsz,'b':True,'c':color,'align':PP_ALIGN.CENTER}])
    text(sl,x,y+1.02,w,0.4,[{'t':label,'sz':12.5,'c':MUTED,'align':PP_ALIGN.CENTER}])

def note(sl, t, y=6.85):
    text(sl,0.75,y,11.8,0.35,[{'t':t,'sz':11,'c':MUTED}])

_HERE = os.path.dirname(os.path.abspath(__file__))    # presentation/deckgen/
_PRES = os.path.dirname(_HERE)                        # presentation/
_REPO = os.path.dirname(_PRES)                        # 저장소 루트
QR_PATH = os.path.join(_REPO, 'demo', 'qr_repo.png')
REPO_URL = 'github.com/giyeop-cody/token-cost-lab'


def qr(sl, x, y, size=1.5, caption=None, csz=11, cw=None):
    """리포 QR. 흰 여백(quiet zone)이 있어야 스캔되므로 흰 판 위에 얹는다."""
    if not os.path.exists(QR_PATH):
        return
    pad = 0.08
    rect(sl, x - pad, y - pad, size + pad * 2, size + pad * 2,
         fill=RGBColor(0xFF, 0xFF, 0xFF), radius=0.05)
    sl.shapes.add_picture(QR_PATH, Inches(x), Inches(y),
                          Inches(size), Inches(size))
    if caption:
        text(sl, x - 0.25, y + size + 0.14, cw or (size + 0.5), 0.3,
             [{'t': caption, 'sz': csz, 'c': MUTED}], align=PP_ALIGN.CENTER)

