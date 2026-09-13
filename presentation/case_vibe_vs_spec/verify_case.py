#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Vibe vs Spec 케이스 — 발표 자료 수치 자동 검증.

usage.json(원자료) 하나만을 진실의 근거로 삼아 모든 파생 수치를 재계산하고,
발표 문서(.md)와 슬라이드 덱(.pptx)에 기재된 값이 그 값과 일치하는지 검사한다.

    python3 presentation/case_vibe_vs_spec/verify_case.py

종료 코드 0 = 전 항목 일치, 1 = 불일치 존재.
"""
from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CASE = ROOT / "presentation" / "case_vibe_vs_spec"
USAGE = ROOT / "demo" / "vibe_vs_spec" / "usage.json"

# 공시 단가 ($ per 1M tokens). 추론은 출력 요율로 과금된다.
PRICES = {
    "GPT-5":            {"input": 1.25, "reasoning": 10.0, "output": 10.0, "cache": 0.125},
    "Claude Sonnet 4.5": {"input": 3.00, "reasoning": 15.0, "output": 15.0, "cache": 0.300},
}
KRW = 1390

results: list[tuple[bool, str, str]] = []


def check(ok: bool, name: str, detail: str = "") -> None:
    results.append((bool(ok), name, detail))


def cost(tokens: dict, price: dict) -> dict:
    return {k: tokens[k] * price[k] / 1e6 for k in price}


def main() -> int:
    usage = json.loads(USAGE.read_text(encoding="utf-8"))
    vibe = usage["runs"]["vibe"]["tokens"]
    spec = usage["runs"]["spec"]["tokens"]

    # ---- 1. 토큰 원자료 무결성 -------------------------------------------
    for label, t in (("vibe", vibe), ("spec", spec)):
        parts = t["input"] + t["reasoning"] + t["output"] + t["cache"]
        check(parts == t["total"], f"[{label}] 합계 = 항목 합",
              f"{parts} vs {t['total']}")

    d_tok = spec["total"] - vibe["total"]
    d_pct = d_tok / vibe["total"] * 100
    check(d_tok == 142, "토큰 증가분 = +142", str(d_tok))
    check(abs(d_pct - 1.8) < 0.05, "토큰 증가율 = +1.8%", f"{d_pct:+.2f}%")

    for key, want in (("input", 137), ("reasoning", -21), ("output", -7), ("cache", 50)):
        got = (spec[key] - vibe[key]) / vibe[key] * 100
        check(abs(got - want) < 1.0, f"{key} 변화율 = {want:+d}%", f"{got:+.1f}%")

    ratio = spec["input"] / vibe["input"]
    check(abs(ratio - 2.37) < 0.01, "입력 배수 = 2.37배", f"{ratio:.2f}배")

    # ---- 2. 비용 재계산 ---------------------------------------------------
    money: dict[str, dict] = {}
    for model, price in PRICES.items():
        cv, cs = cost(vibe, price), cost(spec, price)
        tv, ts = sum(cv.values()), sum(cs.values())
        money[model] = {"vibe": tv, "spec": ts, "delta": (ts - tv) / tv * 100,
                        "share": {k: v / tv * 100 for k, v in cv.items()}}

    g, c = money["GPT-5"], money["Claude Sonnet 4.5"]
    check(abs(g["vibe"] - 0.07135) < 5e-6, "GPT-5 Vibe = $0.07135", f"${g['vibe']:.5f}")
    check(abs(g["spec"] - 0.06540) < 5e-6, "GPT-5 Spec = $0.06540", f"${g['spec']:.5f}")
    check(abs(g["delta"] + 8.3) < 0.06, "GPT-5 절감 = −8.3%", f"{g['delta']:.2f}%")
    check(abs(c["vibe"] - 0.10760) < 5e-6, "Claude Vibe = $0.10760", f"${c['vibe']:.5f}")
    check(abs(c["spec"] - 0.09945) < 5e-6, "Claude Spec = $0.09945", f"${c['spec']:.5f}")
    check(abs(c["delta"] + 7.6) < 0.06, "Claude 절감 = −7.6%", f"{c['delta']:.2f}%")

    check(abs(round(g["vibe"] * KRW, 1) - 99.2) < 0.05, "GPT-5 Vibe 원화 = ₩99.2",
          f"₩{g['vibe']*KRW:.1f}")
    check(abs(round(g["spec"] * KRW, 1) - 90.9) < 0.05, "GPT-5 Spec 원화 = ₩90.9",
          f"₩{g['spec']*KRW:.1f}")

    heavy = g["share"]["reasoning"] + g["share"]["output"]
    check(heavy > 99.0, "GPT-5 추론+출력 비중 > 99%", f"{heavy:.1f}%")
    check(abs(g["share"]["output"] - 81.71) < 0.05, "GPT-5 출력 비중 = 81.71%",
          f"{g['share']['output']:.2f}%")

    # 재생성 1회 vs 스펙 입력 비용 배수
    spec_in = spec["input"] * PRICES["GPT-5"]["input"] / 1e6
    mult = g["spec"] / spec_in
    check(abs(mult - 45) < 1.0, "재생성 1회 = 스펙 입력의 45배", f"{mult:.0f}배")

    # ---- 3. 문서 기재값 대조 ---------------------------------------------
    docs = {p.name: p.read_text(encoding="utf-8") for p in CASE.glob("*.md")}
    joined = "\n".join(docs.values())

    must = ["$0.07135", "$0.06540", "−8.3%", "−7.6%", "+1.8%", "16/16", "2.37배"]
    for token in must:
        check(token in joined, f"문서에 {token} 존재")

    stale = ["$0.07145", "−8.5%", "-8.5%", "₩99.3", "토큰 141개"]
    for token in stale:
        hits = [n for n, t in docs.items() if token in t]
        check(not hits, f"폐기값 '{token}' 미잔존", ", ".join(hits))

    # ---- 4. 슬라이드 덱 대조 ---------------------------------------------
    deck = CASE / "Vibe-vs-Spec_토큰비용_리포트.pptx"
    try:
        from pptx import Presentation
    except ImportError:
        check(True, "pptx 검사 건너뜀 (python-pptx 미설치)")
    else:
        prs = Presentation(deck)
        check(len(prs.slides) == 9, "케이스 덱 = 9장", str(len(prs.slides)))
        text = []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    text.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        text.extend(cell.text for cell in r.cells)
        blob = "\n".join(text)
        for token in ("7,868", "8,010", "+1.8%", "−8.3%", "16/16", "$0.07135", "$0.06540"):
            check(token in blob, f"덱에 {token} 존재")
        check("\ufffd" not in blob, "덱에 깨진 문자(U+FFFD) 없음")
        for token in ("−8.5%", "$0.07145"):
            check(token not in blob, f"덱에 폐기값 '{token}' 없음")

    # ---- 5. 본 덱 장수 정합 ----------------------------------------------
    main_deck = ROOT / "presentation" / "토큰_절약_발표.pptx"
    if main_deck.exists():
        try:
            from pptx import Presentation as P2
            n = len(P2(main_deck).slides)
        except ImportError:
            n = None
        if n is not None:
            check(n == 29, "본 덱 = 29장", str(n))
            pres_docs = "\n".join(
                p.read_text(encoding="utf-8")
                for p in (ROOT / "presentation").glob("*.md"))
            check("28장" not in pres_docs, "문서에 '28장' 오기 없음")

    # ---- 출력 --------------------------------------------------------------
    fails = [r for r in results if not r[0]]
    for ok, name, detail in results:
        mark = "PASS" if ok else "FAIL"
        line = f"[{mark}] {name}"
        if detail:
            line += f"  ({detail})"
        print(line)
    print("-" * 60)
    print(f"{len(results) - len(fails)}/{len(results)} PASS")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
