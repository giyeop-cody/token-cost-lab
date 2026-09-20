#!/usr/bin/env python3
"""Strict artifact verification, including every actual slide's text.

Default: all five published PPTX files must exist. --pptx checks a supplied
artifact against --deck (default token_cost), useful for mutation tests.
--with-pdf also verifies page counts and all slide text in the generated PDFs.
No missing dependency or artifact is counted as PASS.
"""
import argparse
import json
import math
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from lab.evidence import build
from presentation.deckgen.build_verified import (expected_texts, notes_text, script_text,
                                                specifications)


def same_value(a, b):
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(same_value(a[k], b[k]) for k in a if k != "tiktoken_version")
    if isinstance(a, list) and isinstance(b, list):
        return len(a) == len(b) and all(same_value(x, y) for x, y in zip(a, b))
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-12)
    return a == b


def shape_texts(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    texts = []
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(shape_texts(sh.shapes))
        elif getattr(sh, "has_text_frame", False) and sh.text_frame.text:
            texts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False):
            texts.extend(cell.text for row in sh.table.rows for cell in row.cells)
    return texts


def verify_pptx(path, specs):
    from pptx import Presentation
    checks = []
    def add(name, ok, detail=""):
        checks.append(dict(name=name, ok=bool(ok), detail=detail))
    path = Path(path)
    add("artifact exists", path.is_file(), str(path))
    if not path.is_file():
        return checks
    prs = Presentation(path)
    add("slide count", len(prs.slides) == len(specs), f"{len(prs.slides)} vs {len(specs)}")
    for i, (sl, spec) in enumerate(zip(prs.slides, specs), 1):
        actual, expected = shape_texts(sl.shapes), expected_texts(spec, i)
        mismatch = ""
        if actual != expected:
            mismatch = f"actual text differs from regenerated evidence specification (slide {i})"
        add(f"slide {i} all text/figures/qualifiers", actual == expected, mismatch)
        expected_notes = notes_text(spec)
        add(f"slide {i} speaker notes", sl.notes_slide.notes_text_frame.text == expected_notes)
        for sh in sl.shapes:
            if getattr(sh, "has_text_frame", False) and sh.text:
                add(f"slide {i} {sh.name} within slide", sh.left >= 0 and sh.top >= 0 and
                    sh.left+sh.width <= prs.slide_width+100 and sh.top+sh.height <= prs.slide_height+100)
    return checks


def normalize(text):
    return re.sub(r"\s+", "", text)


def verify_pdf(path, specs):
    import pymupdf as fitz
    path = Path(path)
    checks = [dict(name="PDF exists", ok=path.is_file(), detail=str(path))]
    if not path.is_file():
        return checks
    doc = fitz.open(path)
    checks.append(dict(name="PDF page count", ok=len(doc) == len(specs), detail=str(len(doc))))
    for i, (page, spec) in enumerate(zip(doc, specs), 1):
        text = normalize(page.get_text())
        missing = [s for s in expected_texts(spec, i) if normalize(s) not in text]
        exact = text == normalize("\n".join(expected_texts(spec, i)))
        checks.append(dict(name=f"PDF page {i} all text", ok=exact, detail=" | ".join(missing) if missing else ("" if exact else "unexpected/reordered PDF text")))
    doc.close()
    return checks


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--pptx", type=Path)
    ap.add_argument("--deck", default="token_cost")
    ap.add_argument("--with-pdf", action="store_true")
    ap.add_argument("--verbose", action="store_true")
    ap.add_argument("--json", type=Path)
    args = ap.parse_args()
    checks = []
    try:
        evidence = build()
        snapshot = json.loads((ROOT / "results/audit_metrics.json").read_text(encoding="utf-8"))
        checks.append(dict(name="audit_metrics matches primary evidence", ok=same_value(snapshot, evidence), detail=""))
        specs = specifications(evidence)
        selected = {args.deck: args.pptx} if args.pptx else {k: ROOT / "presentation" / (k+".pptx") for k in specs}
        for name, path in selected.items():
            current = verify_pptx(path, specs[name])
            if not args.pptx:
                script_name = {"token_cost":"script", "token_cost_main":"script_main",
                               "token_cost_bonus":"script_bonus", "token_cost_agent":"script_agent",
                               "case_vibe_vs_spec/vibe_vs_spec":"case_vibe_vs_spec/script_case"}[name]
                script_path = ROOT / "presentation" / (script_name + ".md")
                current.append(dict(name="speaker script", ok=script_path.is_file() and script_path.read_text() == script_text(specs[name]), detail=str(script_path)))
            if args.with_pdf:
                current += verify_pdf(path.with_suffix(".pdf"), specs[name])
            for c in current:
                c["name"] = name + ": " + c["name"]
            checks += current
    except Exception as exc:
        checks.append(dict(name="verification setup/runtime", ok=False, detail=f"{type(exc).__name__}: {exc}"))
    for check in checks:
        if args.verbose or not check["ok"]:
            print(f"{'PASS' if check['ok'] else 'FAIL'} {check['name']} {check['detail']}")
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps({"checks": checks}, ensure_ascii=False, indent=2) + "\n")
    failures = sum(not c["ok"] for c in checks)
    print(f"{len(checks)-failures} PASS / {failures} FAIL / 0 SKIP (text/figures/layout bounds only; not external fact-proof)")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
