import json
import os
from pathlib import Path
import subprocess
import sys

import pytest
from pptx import Presentation

from presentation.deckgen.build_verified import specifications
from tools.verify_deck import verify_pptx, verify_pdf

ROOT=Path(__file__).resolve().parents[1]
DECKS=["token_cost","token_cost_main","token_cost_bonus","token_cost_agent","case_vibe_vs_spec/vibe_vs_spec"]


@pytest.mark.parametrize("name",DECKS)
def test_published_pptx_and_pdf_full_text(name,evidence):
    spec=specifications(evidence)[name]
    path=ROOT/"presentation"/(name+".pptx")
    checks=verify_pptx(path,spec)+verify_pdf(path.with_suffix(".pdf"),spec)
    assert all(c["ok"] for c in checks), [c for c in checks if not c["ok"]]


def test_case_deck_restored_figures_are_byte_identical_to_sources():
    """케이스 덱에 다시 넣은 스크린샷이 원본 PNG와 바이트 동일한지 검사.

    덱을 다시 만들면서 그림을 재인코딩·대체하면 잡힌다. (그림 없이 텍스트만 남은 상태도 잡힌다.)
    """
    import hashlib
    import zipfile
    wanted = {f"demo/vibe_vs_spec/shots/{n}.png" for n in
              ("vibe_top", "spec_top", "spec_mobile", "vibe_spec_full_2up")}
    sources = {hashlib.sha256((ROOT / w).read_bytes()).hexdigest() for w in wanted}
    for w in wanted:
        assert (ROOT / w).is_file(), w
    with zipfile.ZipFile(ROOT / "presentation/case_vibe_vs_spec/vibe_vs_spec.pptx") as z:
        embedded = {hashlib.sha256(z.read(m)).hexdigest()
                    for m in z.namelist() if m.startswith("ppt/media/")}
    missing = sources - embedded
    assert not missing, f"케이스 덱에 원본 그대로 박히지 않은 그림 {len(missing)}개"


def test_case_deck_fact_tables_match_usage_and_price_sources(evidence):
    """케이스 덱의 토큰·단가 표 값이 원자료 산술과 일치하는지 검사."""
    usage = json.loads((ROOT / "demo/vibe_vs_spec/usage.json").read_text(encoding="utf-8"))["runs"]
    vibe, spec = usage["vibe"]["tokens"], usage["spec"]["tokens"]
    deck = {s["title"]: s for s in specifications(evidence)["case_vibe_vs_spec/vibe_vs_spec"]}
    facts = {row[0]: row[1] for s in deck.values() for row in s["facts"]}
    assert facts["비캐시 입력"] == f"{vibe['input']:,} → {spec['input']:,}  (증가 {spec['input']-vibe['input']})"
    assert facts["추론+보이는 출력"] == (f"{vibe['reasoning']+vibe['output']:,} → "
                                        f"{spec['reasoning']+spec['output']:,}  "
                                        f"(감소 {vibe['reasoning']+vibe['output']-spec['reasoning']-spec['output']})")
    assert facts["합계"] == f"{vibe['total']:,} → {spec['total']:,}  (증가 1.8%)"
    gpt = deck["세 단가로 같은 추정치를 환산"]["facts"][0][1]
    want = usage["vibe"]["cost_usd"]["total"], usage["spec"]["cost_usd"]["total"]
    assert f"${want[0]:.7f} → ${want[1]:.7f}" in gpt, gpt
    assert deck["단가비가 손익을 결정한다"]["facts"][1][1] == "r = 1.0024"


def test_240_to_999_mutation_fails_actual_deck_and_cli(tmp_path,evidence):
    prs=Presentation(ROOT/"presentation/token_cost.pptx")
    edits=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if "240" in run.text:
                            edits += run.text.count("240")
                            run.text=run.text.replace("240","999")
    assert edits > 0
    path=tmp_path/"mutated.pptx"
    prs.save(path)
    checks=verify_pptx(path,specifications(evidence)["token_cost"])
    assert any(not c["ok"] and "all text" in c["name"] for c in checks)
    proc=subprocess.run([sys.executable,str(ROOT/"tools/verify_deck.py"),"--pptx",str(path)],
                        cwd=ROOT,text=True,capture_output=True,timeout=60)
    assert proc.returncode == 1,proc.stdout+proc.stderr
    assert "FAIL" in proc.stdout


def test_missing_pptx_is_a_failure(tmp_path,evidence):
    checks=verify_pptx(tmp_path/"missing.pptx",specifications(evidence)["token_cost"])
    assert checks and not checks[0]["ok"]


@pytest.mark.parametrize("name",["test_router.py","test_orchestrator.py","test_intent_guard.py","test_lean.py","test_mcp.py"])
def test_legacy_checks_are_executable_and_pass(name):
    proc=subprocess.run([sys.executable,str(ROOT/"agent_setup"/name)],cwd=ROOT,
                        text=True,capture_output=True,timeout=90)
    assert proc.returncode == 0,proc.stdout+proc.stderr
    assert "FAIL" not in proc.stdout or "0 FAIL" in proc.stdout


@pytest.mark.parametrize("name",["test_intent_guard.py","test_orchestrator.py"])
def test_forced_assertion_failure_exits_nonzero(name,tmp_path):
    original=(ROOT/"agent_setup"/name).read_text()
    assert "if not c:" in original
    mutated=original.replace("if not c:","if True:  # intentional regression mutation")
    path=tmp_path/name
    path.write_text(mutated)
    env=dict(os.environ,PYTHONPATH=str(ROOT/"agent_setup")+os.pathsep+str(ROOT))
    proc=subprocess.run([sys.executable,str(path)],cwd=ROOT,env=env,
                        text=True,capture_output=True,timeout=45)
    assert proc.returncode == 1,proc.stdout+proc.stderr
    assert "FAIL" in proc.stdout


@pytest.mark.parametrize("first_fails",[True,False])
def test_shell_preserves_pipeline_failure_and_separates_skips(tmp_path,first_fails):
    # A fake Python child lets us test pipefail without MCP/Inspector network installs.
    bin_dir=tmp_path/"bin"
    bin_dir.mkdir()
    python=bin_dir/"python3"
    python.write_text('#!/bin/sh\nif [ "$1" = "-c" ]; then exit 1; fi\necho "controlled child"\nexit '+('1' if first_fails else '0')+'\n')
    python.chmod(0o755)
    for cmd in ("tail","dirname"):
        (bin_dir/cmd).symlink_to("/usr/bin/"+cmd)
    proc=subprocess.run(["/bin/bash",str(ROOT/"agent_setup/verify_all.sh")],cwd=ROOT,
                        env=dict(os.environ,PATH=str(bin_dir)),capture_output=True,text=True,timeout=30)
    assert proc.returncode == (1 if first_fails else 2),proc.stdout+proc.stderr
    if not first_fails:
        assert "미검증 (전체 통과 아님)" in proc.stdout
        assert "네 층위 프로토콜 연결 통과" not in proc.stdout


def test_migrated_paths_exist_and_filenames_are_utf8():
    rows=json.loads((ROOT/"docs/path_migration.json").read_text())
    assert len(rows) == 27
    for row in rows:
        assert (ROOT/row["new_path"]).exists(),row
    for p in ROOT.rglob("*"):
        if ".git" not in p.parts and "__pycache__" not in p.parts:
            str(p.relative_to(ROOT)).encode("utf-8",errors="strict")


def test_real_ci_workflow_exists():
    content=(ROOT/".github/workflows/ci.yml").read_text()
    assert "pytest" in content and "playwright" in content and "verify_deck" in content


def test_extra_wrong_pdf_text_is_not_ignored(tmp_path,evidence):
    import pymupdf
    source=ROOT/"presentation/token_cost.pdf"
    d=pymupdf.open(source)
    d[0].insert_text((40,40),"999",fontsize=12)
    path=tmp_path/"wrong-extra-text.pdf"
    d.save(path)
    d.close()
    checks=verify_pdf(path,specifications(evidence)["token_cost"])
    assert any(not c["ok"] and "page 1 " in c["name"] for c in checks)


@pytest.mark.parametrize("number",["08","10","11"])
def test_run_all_live_without_keys_is_unverified(number,monkeypatch,capsys):
    import run_all
    from types import SimpleNamespace
    monkeypatch.setattr(sys,"argv",["run_all","--live","--only",number])
    monkeypatch.setattr(run_all.subprocess,"run",lambda *a,**kw:SimpleNamespace(returncode=0))
    assert run_all.main() == 2
    text=capsys.readouterr().out
    assert "미검증" in text and "모든 실험 완료" not in text


def test_run_all_default_never_enables_live(monkeypatch):
    import run_all
    from types import SimpleNamespace
    commands=[]
    def run(cmd,**kw):
        commands.append(cmd)
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(sys,"argv",["run_all"])
    monkeypatch.setattr(run_all.subprocess,"run",run)
    assert run_all.main() in (None,0)
    assert len(commands) == 11
    assert all("--live" not in c for c in commands)
    for number in ("exp08","exp10"):
        assert all("--dry-run" in c for c in commands if number in c[1])


def test_inspector_output_cannot_hide_nonzero_status(tmp_path):
    bin_dir=tmp_path/"bin"
    bin_dir.mkdir()
    py=bin_dir/"python3"
    py.write_text('#!/bin/sh\nif [ "$1" = "-c" ]; then case "$2" in *"import mcp"*) exit 0;; *) echo mid; exit 0;; esac; fi\necho "controlled child"\nexit 0\n')
    py.chmod(0o755)
    npx=bin_dir/"npx"
    npx.write_text('#!/bin/sh\nprintf \'"name":1\\n"name":2\\n"name":3\\n"name":4\\n"name":5\\n"name":6\\n\'\nexit 1\n')
    npx.chmod(0o755)
    for cmd in ("tail","dirname","grep"):
        (bin_dir/cmd).symlink_to("/usr/bin/"+cmd)
    p=subprocess.run(["/bin/bash",str(ROOT/"agent_setup/verify_all.sh")],cwd=ROOT,
                     env=dict(os.environ,PATH=str(bin_dir)),text=True,capture_output=True,timeout=30)
    assert p.returncode == 1,p.stdout+p.stderr
    assert "PASS  Inspector" not in p.stdout


AUDIT_ONLY_TOKENS = ("반례", "미측정", "철회", "격리", "비일치", "정정 이력", "감사", "한계")


def test_deck_slides_keep_messages_and_notes_carry_scope(evidence):
    """발표체 규칙을 고정한다: 슬라이드는 메시지·숫자, 검증 범위·한계는 발표자 노트.

    2026-09-21 지적(“발표 자료를 만드는 거지, 실험 결과와 방향성 반례를 적는 것이 아님”)에 따라
    슬라이드 본문에서 한계·반례 진술을 걷어냈다. 다시 슬라이드로 올라오면 이 검사가 잡는다.
    """
    from presentation.deckgen.build_verified import notes_text
    decks = specifications(evidence)
    assert set(decks) == set(DECKS)
    for name, specs in decks.items():
        for idx, spec in enumerate(specs, 1):
            face = "\n".join([spec["title"], spec["hero"], *spec["lines"],
                              *[cell for row in spec["facts"] for cell in row],
                              *[caption for _, caption in spec["images"]]])
            for token in AUDIT_ONLY_TOKENS:
                assert token not in face, f"{name} {idx}장 슬라이드에 검증 서술: {token}"
            assert len(spec["script"]) >= 2, f"{name} {idx}장 발표 멘트 부족"
            notes = notes_text(spec)
            assert notes.startswith(spec["script"][0]), f"{name} {idx}장 노트 첫 문단"
            assert f"증거 범위: {spec['scope']}" in notes, f"{name} {idx}장 증거 범위 누락"
            assert f"근거: {spec['source']}" in notes, f"{name} {idx}장 근거 누락"


def test_deck_scripts_declare_target_duration_and_per_slide_talk():
    """대본(.md)은 30분 목표 분량과 장별 멘트를 담는다(덱과 같은 원천에서 생성)."""
    from presentation.deckgen.build_verified import script_text
    decks = specifications(json.loads((ROOT / "results/audit_metrics.json").read_text(encoding="utf-8")))
    for name, specs in decks.items():
        text = script_text(specs)
        assert f"{len(specs)}장 · 약 30분" in text, name
        assert text.count("\n\n**멘트**\n\n") == len(specs), name
        for spec in specs:
            assert spec["script"][0] in text, (name, spec["title"])
